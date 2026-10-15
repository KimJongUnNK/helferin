#!/usr/bin/env python3
"""Erzeugt .docx-, .pdf- und .pptx-Dateien für Helferlein.

Liest ein JSON-Objekt von stdin und gibt den Pfad der erstellten Datei auf stdout aus.

  {"kind": "docx", "title": "...", "sections": [{"heading": "...", "text": "..."}], "out_dir": "..."}
  {"kind": "pdf",  "title": "...", "sections": [{"heading": "...", "text": "..."}], "out_dir": "..."}
  {"kind": "pptx", "title": "...", "slides": [{"title": "...", "bullets": ["..."]}], "out_dir": "..."}
"""
import datetime
import html
import json
import os
import re
import sys

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

BULLET_RE = re.compile(r"^\s*(?:[-•*–]|\d+[.)])\s+")


def safe_filename(title: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\n\r\t]+', " ", title or "").strip()
    name = re.sub(r"\s+", " ", name)[:80].strip(" .")
    return name or "Dokument"


def unique_path(out_dir: str, base: str, ext: str) -> str:
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{base}.{ext}")
    n = 2
    while os.path.exists(path):
        path = os.path.join(out_dir, f"{base} {n}.{ext}")
        n += 1
    return path


def add_text_block(doc, text: str) -> None:
    """Absätze durch Leerzeilen, Listenpunkte mit '- ' oder '1.'; einzelne Zeilenumbrüche bleiben erhalten."""
    for block in re.split(r"\n\s*\n", (text or "").strip()):
        lines = [l for l in block.split("\n") if l.strip()]
        if not lines:
            continue
        if all(BULLET_RE.match(l) for l in lines):
            numbered = all(re.match(r"^\s*\d+[.)]\s+", l) for l in lines)
            style = "List Number" if numbered else "List Bullet"
            for l in lines:
                doc.add_paragraph(BULLET_RE.sub("", l).strip(), style=style)
            continue
        para = doc.add_paragraph()
        for i, line in enumerate(lines):
            run = para.add_run(line.strip())
            if i < len(lines) - 1:
                run.add_break()


def build_docx(payload: dict) -> str:
    title = (payload.get("title") or "Dokument").strip()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)

    sections = payload.get("sections") or []
    # Briefe haben keine Überschriften – dann auch keinen großen Titel drüber
    is_letter = sections and all(not (s.get("heading") or "").strip() for s in sections)
    if not is_letter:
        doc.add_heading(title, level=0)
    for s in sections:
        heading = (s.get("heading") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        add_text_block(doc, s.get("text") or "")

    path = unique_path(payload["out_dir"], safe_filename(title), "docx")
    doc.save(path)
    return path


def split_blocks(text: str):
    """Liefert ("bullets"|"numbers"|"para", [zeilen]) je Absatz – gleiche Regeln wie bei docx."""
    for block in re.split(r"\n\s*\n", (text or "").strip()):
        lines = [l for l in block.split("\n") if l.strip()]
        if not lines:
            continue
        if all(BULLET_RE.match(l) for l in lines):
            numbered = all(re.match(r"^\s*\d+[.)]\s+", l) for l in lines)
            yield ("numbers" if numbered else "bullets"), [BULLET_RE.sub("", l).strip() for l in lines]
        else:
            yield "para", [l.strip() for l in lines]


def build_pdf(payload: dict) -> str:
    title = (payload.get("title") or "Dokument").strip()
    sections = payload.get("sections") or []
    is_letter = sections and all(not (s.get("heading") or "").strip() for s in sections)

    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontSize, body.leading, body.spaceAfter = 11.5, 16, 8
    esc = lambda t: html.escape(t, quote=False)

    story = []
    if not is_letter:
        story += [Paragraph(esc(title), styles["Title"]), Spacer(1, 0.3 * cm)]
    for s in sections:
        heading = (s.get("heading") or "").strip()
        if heading:
            story.append(Paragraph(esc(heading), styles["Heading2"]))
        for kind, lines in split_blocks(s.get("text") or ""):
            if kind == "para":
                story.append(Paragraph("<br/>".join(esc(l) for l in lines), body))
            else:
                items = [ListItem(Paragraph(esc(l), body)) for l in lines]
                story.append(ListFlowable(items, bulletType="1" if kind == "numbers" else "bullet", leftIndent=14))

    path = unique_path(payload["out_dir"], safe_filename(title), "pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, title=title,
                            leftMargin=2.5 * cm, rightMargin=2.2 * cm, topMargin=2.2 * cm, bottomMargin=2.2 * cm)
    doc.build(story or [Spacer(1, 1)])
    return path


def build_pptx(payload: dict) -> str:
    title = (payload.get("title") or "Präsentation").strip()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = datetime.date.today().strftime("%d.%m.%Y")

    for s in payload.get("slides") or []:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = (s.get("title") or "").strip()
        body = slide.placeholders[1]
        body.width = prs.slide_width - Inches(1.0)
        tf = body.text_frame
        bullets = [BULLET_RE.sub("", b).strip() for b in (s.get("bullets") or []) if b and b.strip()]
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = PptPt(24)

    path = unique_path(payload["out_dir"], safe_filename(title), "pptx")
    prs.save(path)
    return path


def main() -> int:
    try:
        payload = json.load(sys.stdin)
        if not payload.get("out_dir"):
            raise ValueError("out_dir fehlt")
        kind = payload.get("kind")
        if kind == "docx":
            path = build_docx(payload)
        elif kind == "pdf":
            path = build_pdf(payload)
        elif kind == "pptx":
            path = build_pptx(payload)
        else:
            raise ValueError(f"Unbekannter Typ: {kind!r}")
    except Exception as e:  # Fehler gehen als Text an die App zurück
        print(f"{type(e).__name__}: {e}", file=sys.stderr)
        return 1
    print(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
