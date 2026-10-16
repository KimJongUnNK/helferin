import json
import os
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.lang import MSO_LANGUAGE_ID

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "craftdocs.py")


def run(payload):
    p = subprocess.run([sys.executable, SCRIPT], input=json.dumps(payload), capture_output=True, text=True)
    return p.returncode, p.stdout.strip(), p.stderr.strip()


class CraftDocsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def test_docx_with_headings_bullets_and_numbers(self):
        code, path, err = run({
            "kind": "docx", "title": "Einkaufsliste: Wochenende", "out_dir": self.tmp,
            "sections": [
                {"heading": "Obst", "text": "- Äpfel\n- Bananen"},
                {"heading": "Ablauf", "text": "1. Einkaufen\n2. Kochen"},
                {"heading": "Notiz", "text": "Erster Absatz.\n\nZweiter Absatz\nmit Umbruch."},
            ]})
        self.assertEqual(code, 0, err)
        self.assertTrue(path.endswith("Einkaufsliste Wochenende.docx"))
        doc = Document(path)
        texts = [(p.style.name, p.text) for p in doc.paragraphs]
        self.assertIn(("Title", "Einkaufsliste: Wochenende"), texts)
        self.assertIn(("List Bullet", "Äpfel"), texts)
        self.assertIn(("List Number", "Kochen"), texts)
        self.assertIn(("Normal", "Zweiter Absatz\nmit Umbruch."), texts)

    def test_docx_language_is_german(self):
        code, path, err = run({"kind": "docx", "title": "Sprache", "lang": "de-AT", "out_dir": self.tmp,
                               "sections": [{"heading": "Hallo", "text": "Grüß Gott"}]})
        self.assertEqual(code, 0, err)
        import zipfile
        with zipfile.ZipFile(path) as z:
            styles = z.read("word/styles.xml").decode()
        self.assertIn('w:val="de-AT"', styles)
        self.assertNotIn('w:val="en-US"', styles)
        code, path, err = run({"kind": "docx", "title": "Standard", "out_dir": self.tmp, "sections": []})
        with zipfile.ZipFile(path) as z:
            self.assertIn('w:val="de-DE"', z.read("word/styles.xml").decode())
        code, path, err = run({"kind": "docx", "title": "Kaputt", "lang": "<x>", "out_dir": self.tmp, "sections": []})
        self.assertEqual(code, 0, err)

    def test_font_settings(self):
        code, path, err = run({"kind": "docx", "title": "Schrift", "font": "Georgia", "font_size": 14,
                               "out_dir": self.tmp, "sections": [{"heading": "", "text": "Hallo"}]})
        self.assertEqual(code, 0, err)
        normal = Document(path).styles["Normal"].font
        self.assertEqual((normal.name, normal.size.pt), ("Georgia", 14))
        code, path, err = run({"kind": "pdf", "title": "Schrift", "font": "Times New Roman", "font_size": 99,
                               "out_dir": self.tmp, "sections": [{"heading": "", "text": "Hallo"}]})
        self.assertEqual(code, 0, err)
        with open(path, "rb") as f:
            self.assertIn(b"Times-Roman", f.read())

    def test_letter_has_no_title_heading(self):
        code, path, err = run({
            "kind": "docx", "title": "Brief an Tante Gerti", "out_dir": self.tmp,
            "sections": [{"heading": "", "text": "Liebe Gerti,"}, {"heading": "", "text": "Viele Grüße\nMichaela"}]})
        self.assertEqual(code, 0, err)
        styles = {p.style.name for p in Document(path).paragraphs}
        self.assertNotIn("Title", styles)

    def test_no_overwrite(self):
        payload = {"kind": "docx", "title": "Test", "out_dir": self.tmp, "sections": []}
        _, a, _ = run(payload)
        _, b, _ = run(payload)
        self.assertNotEqual(a, b)
        self.assertTrue(b.endswith("Test 2.docx"))

    def test_pptx(self):
        code, path, err = run({
            "kind": "pptx", "title": "Gartenclub", "out_dir": self.tmp,
            "slides": [{"title": "Termine", "bullets": ["- März: Schnitt", "April: Säen", ""]},
                       {"title": "Kosten", "bullets": ["20 € pro Jahr"]}]})
        self.assertEqual(code, 0, err)
        prs = Presentation(path)
        self.assertEqual(len(prs.slides), 3)
        langs = {r.font.language_id for s in prs.slides for sh in s.shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs}
        self.assertEqual(langs, {MSO_LANGUAGE_ID.GERMAN})
        body = prs.slides[1].placeholders[1].text_frame
        self.assertEqual([p.text for p in body.paragraphs], ["März: Schnitt", "April: Säen"])

    def test_pdf_with_umlauts_lists_and_escaping(self):
        code, path, err = run({
            "kind": "pdf", "title": "Rezept: Käsespätzle", "out_dir": self.tmp,
            "sections": [{"heading": "Zutaten", "text": "- 500 g Mehl\n- Bergkäse & Zwiebeln"},
                         {"heading": "Zubereitung", "text": "1. Teig rühren\n2. Hobeln <schnell>\n\nGuten Appetit, 5 €!"}]})
        self.assertEqual(code, 0, err)
        self.assertTrue(path.endswith("Rezept Käsespätzle.pdf"))
        with open(path, "rb") as f:
            data = f.read()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 1500)

    def test_pdf_letter_and_empty(self):
        code, _, err = run({"kind": "pdf", "title": "Brief", "out_dir": self.tmp,
                            "sections": [{"heading": "", "text": "Liebe Gerti,"}]})
        self.assertEqual(code, 0, err)
        code, _, err = run({"kind": "pdf", "title": "Leer", "out_dir": self.tmp, "sections": []})
        self.assertEqual(code, 0, err)

    def test_xlsx_numbers_formulas_header(self):
        code, path, err = run({
            "kind": "xlsx", "title": "Haushaltsbuch: Oktober", "out_dir": self.tmp,
            "columns": ["Posten", "Betrag", "Anteil"],
            "rows": [["Miete", "750", "50%"], ["Strom", "1.234,50", ""], ["Kaffee", "12,5 €", ""],
                     ["Summe", "=SUM(B2:B4)", ""]]})
        self.assertEqual(code, 0, err)
        self.assertTrue(path.endswith("Haushaltsbuch Oktober.xlsx"))
        ws = load_workbook(path).active
        self.assertEqual(ws.title, "Haushaltsbuch  Oktober")
        self.assertEqual(ws["A1"].value, "Posten")
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["B2"].value, 750)
        self.assertEqual(ws["C2"].value, 0.5)
        self.assertEqual(ws["B3"].value, 1234.5)
        self.assertEqual(ws["B4"].value, 12.5)
        self.assertEqual(ws["B5"].value, "=SUM(B2:B4)")
        self.assertEqual(ws.freeze_panes, "A2")

    def test_xlsx_empty(self):
        code, path, err = run({"kind": "xlsx", "title": "", "out_dir": self.tmp})
        self.assertEqual(code, 0, err)
        self.assertTrue(path.endswith("Tabelle.xlsx"))

    def test_bad_input_returns_error(self):
        code, _, err = run({"kind": "odt", "out_dir": self.tmp})
        self.assertEqual(code, 1)
        self.assertIn("Unbekannter Typ", err)
        code, _, err = run({"kind": "docx"})
        self.assertEqual(code, 1)


if __name__ == "__main__":
    unittest.main()
